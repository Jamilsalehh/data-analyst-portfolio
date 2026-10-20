"""
Builds the Excel workbook the Power Automate flow writes into:
  templates/QA_Audit_Log_TEMPLATE.xlsx  - empty log with one example row (what you upload to OneDrive)
  data/QA_Audit_Log_SAMPLE.xlsx         - the same workbook filled with the 6 months of synthetic audits
Scores, bands and feedback times are Excel formulas, not pasted values.
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parents[1]
D = ROOT / "data"
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); B = Font(name=F, size=10)
HDR = PatternFill("solid", fgColor="1F4E79"); CALC = PatternFill("solid", fgColor="DDEBF7")

audits = pd.read_csv(D / "audits.csv", parse_dates=["audit_ts", "feedback_sent_at", "agent_acknowledged_at"])
agents = pd.read_csv(D / "agents.csv")
rubric = pd.read_csv(D / "rubric.csv")
bands = pd.read_csv(D / "score_bands.csv").sort_values("min_score")
audits = audits.merge(agents[["agent_id", "agent_name", "team"]], on="agent_id")

COLS = ["audit_id", "audit_ts", "reviewer", "agent_id", "agent_name", "team", "channel", "interaction_ref",
        "c1_score", "c2_score", "c3_score", "c4_score", "c5_score", "c6_score", "critical_error", "flagged_issues",
        "feedback_sent_at", "agent_acknowledged_at", "process", "score", "band", "hours_to_feedback"]
W = "Rubric!$C$"
SCORE = ('=ROUND(MIN(IF(OR(O{r}=TRUE,O{r}="TRUE",O{r}="Yes"),59,100),'
         "(I{r}*" + W + "2+J{r}*" + W + "3+K{r}*" + W + "4+L{r}*" + W + "5+M{r}*" + W +
         "6+N{r}*" + W + "7)/5),1)")
BAND = "=LOOKUP(T{r},Bands!$A$2:$A$6,Bands!$B$2:$B$6)"
HOURS = '=IF(Q{r}="","",ROUND((Q{r}-B{r})*24,1))'


def build(rows, path, example=False):
    wb = Workbook()
    ws = wb.active; ws.title = "Audits"
    for c, h in enumerate(COLS, 1):
        x = ws.cell(1, c, h); x.font = H; x.fill = HDR
    for i, r in enumerate(rows.itertuples(index=False), start=2):
        vals = [r.audit_id, r.audit_ts.to_pydatetime(), r.reviewer, r.agent_id, r.agent_name, r.team, r.channel,
                r.interaction_ref, r.c1_score, r.c2_score, r.c3_score, r.c4_score, r.c5_score, r.c6_score,
                bool(r.critical_error), r.flagged_issues if isinstance(r.flagged_issues, str) else "",
                r.feedback_sent_at.to_pydatetime(),
                r.agent_acknowledged_at.to_pydatetime() if pd.notna(r.agent_acknowledged_at) else None, r.process,
                SCORE.format(r=i), BAND.format(r=i), HOURS.format(r=i)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v); cell.font = B
            if c in (2, 17, 18): cell.number_format = "yyyy-mm-dd hh:mm"
            if c >= 20: cell.fill = CALC
    last = 1 + len(rows)
    t = Table(displayName="Audits", ref=f"A1:{chr(64 + len(COLS))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    widths = [11, 16, 14, 9, 18, 14, 8, 12] + [8] * 6 + [9, 40, 16, 16, 30, 7, 18, 10]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    ws.freeze_panes = "B2"
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"I2:N{max(last, 500)}")
    ws.conditional_formatting.add(f"T2:T{last}", CellIsRule(operator="lessThan", formula=["70"], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"T2:T{last}", CellIsRule(operator="greaterThanOrEqual", formula=["90"], fill=PatternFill("solid", fgColor="C6EFCE")))

    rb = wb.create_sheet("Rubric")
    for c, h in enumerate(["criterion_id", "criterion", "weight", "what_it_measures", "how to score (0-5)"], 1):
        x = rb.cell(1, c, h); x.font = H; x.fill = HDR
    for i, r in enumerate(rubric.itertuples(index=False), start=2):
        for c, v in enumerate([r.criterion_id, r.criterion, r.weight, r.what_it_measures,
                               "5 = flawless · 3 = minor gaps · 1 = major gap · 0 = not done"], 1):
            rb.cell(i, c, v).font = B
    rb.cell(8, 2, "Total").font = Font(name=F, bold=True)
    rb.cell(8, 3, "=SUM(C2:C7)").font = Font(name=F, bold=True)
    rb.cell(10, 1, "Score = sum(criterion score / 5 x weight). A critical security error (identity not verified) caps the audit at 59.").font = Font(name=F, italic=True, size=9)
    for c, w in zip("ABCDE", [12, 26, 8, 60, 50]):
        rb.column_dimensions[c].width = w

    bd = wb.create_sheet("Bands")
    for c, h in enumerate(["min_score", "band"], 1):
        x = bd.cell(1, c, h); x.font = H; x.fill = HDR
    for i, r in enumerate(bands.itertuples(index=False), start=2):
        bd.cell(i, 1, int(r.min_score)).font = B; bd.cell(i, 2, r.band).font = B
    bd.cell(8, 1, "Sorted ascending on purpose: the Audits!band formula uses LOOKUP.").font = Font(name=F, italic=True, size=9)
    bd.column_dimensions["B"].width = 22

    ag = wb.create_sheet("Agents")
    hdr = ["agent_id", "agent_name", "team", "team_lead", "agent_email (fill in)", "lead_email (fill in)",
           "audits", "avg_score", "below_70", "critical_errors"]
    for c, h in enumerate(hdr, 1):
        x = ag.cell(1, c, h); x.font = H; x.fill = HDR
    for i, r in enumerate(agents.itertuples(index=False), start=2):
        vals = [r.agent_id, r.agent_name, r.team, r.team_lead, "", "",
                f"=COUNTIF(Audits!$D:$D,A{i})", f'=IFERROR(ROUND(AVERAGEIF(Audits!$D:$D,A{i},Audits!$T:$T),1),"")',
                f'=COUNTIFS(Audits!$D:$D,A{i},Audits!$T:$T,"<70")', f"=COUNTIFS(Audits!$D:$D,A{i},Audits!$O:$O,TRUE)"]
        for c, v in enumerate(vals, 1):
            cell = ag.cell(i, c, v); cell.font = B
            if c >= 7: cell.fill = CALC
    at = Table(displayName="Agents", ref=f"A1:J{1 + len(agents)}")
    at.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ag.add_table(at)
    for c, w in zip("ABCDEFGHIJ", [9, 18, 14, 20, 24, 24, 8, 10, 9, 14]):
        ag.column_dimensions[c].width = w

    hw = wb.create_sheet("How to use", 0)
    notes = [
        ("QA Audit Log", True),
        ("The Power Automate flow adds one row to the Audits table for every Microsoft Forms submission.", False),
        ("White columns are inputs (written by the flow). Blue columns are formulas: do not type in them.", False),
        ("Rubric: change a weight there and every score updates. Weights must total 100.", False),
        ("Bands: score thresholds for the band label.", False),
        ("Agents: fill in agent_email and lead_email. The flow looks these up to send feedback.", False),
        ("Keep the table names 'Audits' and 'Agents': the flow refers to them by name.", False),
    ]
    if example:
        notes.append(("Row 2 of Audits is an example of a completed audit. Delete it before going live.", False))
    for i, (t_, bold) in enumerate(notes, start=1):
        hw.cell(i, 1, t_).font = Font(name=F, bold=bold, size=13 if bold else 10)
    hw.column_dimensions["A"].width = 100
    wb.save(path)
    print("saved", path)


(ROOT / "templates").mkdir(exist_ok=True)
build(audits.head(1), ROOT / "templates/QA_Audit_Log_TEMPLATE.xlsx", example=True)
build(audits, D / "QA_Audit_Log_SAMPLE.xlsx")
